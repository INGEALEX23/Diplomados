import streamlit as st
import pandas as pd
from datetime import date
import io
import os
from streamlit_gsheets import GSheetsConnection

# ReportLab para PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# openpyxl para Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
try:
    from openpyxl.drawing.image import Image as OpenpyxlImage
except ImportError:
    OpenpyxlImage = None

st.set_page_config(
    page_title="Control Escolar - Posgrado e Ingeniería UAdeC",
    layout="wide",
    page_icon="🎓"
)

PATH_LOGO_UADEC = "logo_uadec.png"
PATH_LOGO_POSGRADO = "logo_posgrado.png"

# Conexión persistente a Google Sheets
conn = st.connection("gsheets", type=GSheetsConnection)

def get_data(worksheet_name):
    try:
        df = conn.read(worksheet=worksheet_name, ttl=0)
        return df.dropna(how="all") if df is not None else pd.DataFrame()
    except Exception:
        return pd.DataFrame()

def write_data(worksheet_name, df):
    conn.update(worksheet=worksheet_name, data=df)

# ==========================================
# GESTIÓN DE DATOS EN GOOGLE SHEETS
# ==========================================
def get_grupos():
    df = get_data("GRUPOS")
    if df.empty or "id" not in df.columns:
        return pd.DataFrame(columns=["id", "diplomado", "grupo", "docente"])
    df["id"] = df["id"].astype(int)
    return df

def add_grupo(diplomado, grupo, docente):
    df = get_grupos()
    nuevo_id = int(df["id"].max() + 1) if not df.empty else 1
    nuevo_reg = pd.DataFrame([{
        "id": nuevo_id,
        "diplomado": diplomado.strip().upper(),
        "grupo": grupo.strip().upper(),
        "docente": docente.strip().upper()
    }])
    df = pd.concat([df, nuevo_reg], ignore_index=True)
    write_data("GRUPOS", df)
    return nuevo_id

def delete_grupo(grupo_id):
    df_g = get_grupos()
    write_data("GRUPOS", df_g[df_g["id"] != grupo_id])

    df_a = get_data("ALUMNOS")
    if not df_a.empty and "grupo_id" in df_a.columns:
        write_data("ALUMNOS", df_a[df_a["grupo_id"] != grupo_id])

    df_h = get_data("BITACORA_HORAS")
    if not df_h.empty and "grupo_id" in df_h.columns:
        write_data("BITACORA_HORAS", df_h[df_h["grupo_id"] != grupo_id])

    df_as = get_data("ASISTENCIAS")
    if not df_as.empty and "grupo_id" in df_as.columns:
        write_data("ASISTENCIAS", df_as[df_as["grupo_id"] != grupo_id])

def get_alumnos_df(grupo_id):
    df = get_data("ALUMNOS")
    if df.empty or "grupo_id" not in df.columns:
        return pd.DataFrame(columns=["matricula", "grupo_id", "nombre", "carrera", "telefono", "pago_1", "pago_2", "pago_completo"])
    
    df_g = df[df["grupo_id"] == grupo_id].copy()
    if not df_g.empty:
        df_g["pago_1"] = df_g["pago_1"].astype(bool)
        df_g["pago_2"] = df_g["pago_2"].astype(bool)
        df_g["pago_completo"] = df_g["pago_completo"].astype(bool)
        df_g["matricula"] = df_g["matricula"].astype(str)
        df_g.sort_values(by="nombre", inplace=True)
    return df_g

def save_alumno(matricula, grupo_id, nombre, carrera, telefono, p1, p2, p_comp):
    df = get_data("ALUMNOS")
    if df.empty or "matricula" not in df.columns:
        df = pd.DataFrame(columns=["matricula", "grupo_id", "nombre", "carrera", "telefono", "pago_1", "pago_2", "pago_completo"])
    
    mat_str = str(matricula).strip().upper()
    df = df[~((df["matricula"].astype(str) == mat_str) & (df["grupo_id"] == grupo_id))]

    nuevo = pd.DataFrame([{
        "matricula": mat_str,
        "grupo_id": int(grupo_id),
        "nombre": str(nombre).strip().upper(),
        "carrera": str(carrera).strip().upper(),
        "telefono": str(telefono).strip().upper(),
        "pago_1": bool(p1),
        "pago_2": bool(p2),
        "pago_completo": bool(p_comp)
    }])
    df = pd.concat([df, nuevo], ignore_index=True)
    write_data("ALUMNOS", df)

def delete_alumno(matricula, grupo_id):
    df = get_data("ALUMNOS")
    if not df.empty:
        df = df[~((df["matricula"].astype(str) == str(matricula)) & (df["grupo_id"] == grupo_id))]
        write_data("ALUMNOS", df)

def save_hora(grupo_id, fecha, horas, tema):
    df = get_data("BITACORA_HORAS")
    if df.empty or "id" not in df.columns:
        df = pd.DataFrame(columns=["id", "grupo_id", "fecha", "horas", "tema"])
    n_id = int(df["id"].max() + 1) if not df.empty else 1
    nuevo = pd.DataFrame([{
        "id": n_id,
        "grupo_id": int(grupo_id),
        "fecha": str(fecha),
        "horas": float(horas),
        "tema": str(tema).strip().upper()
    }])
    df = pd.concat([df, nuevo], ignore_index=True)
    write_data("BITACORA_HORAS", df)

def get_horas_df(grupo_id):
    df = get_data("BITACORA_HORAS")
    if df.empty or "grupo_id" not in df.columns:
        return pd.DataFrame(columns=["id", "grupo_id", "fecha", "horas", "tema"])
    df_g = df[df["grupo_id"] == grupo_id].copy()
    if not df_g.empty:
        df_g.sort_values(by="fecha", ascending=False, inplace=True)
    return df_g

def save_asistencia_fecha(grupo_id, fecha, dict_asistencia):
    df = get_data("ASISTENCIAS")
    if df.empty or "fecha" not in df.columns:
        df = pd.DataFrame(columns=["grupo_id", "fecha", "matricula", "presente"])
    
    df = df[~((df["grupo_id"] == grupo_id) & (df["fecha"] == str(fecha)))]
    nuevas = []
    for mat, pres in dict_asistencia.items():
        nuevas.append({
            "grupo_id": int(grupo_id),
            "fecha": str(fecha),
            "matricula": str(mat),
            "presente": bool(pres)
        })
    if nuevas:
        df = pd.concat([df, pd.DataFrame(nuevas)], ignore_index=True)
        write_data("ASISTENCIAS", df)

def get_asistencia_fecha(grupo_id, fecha):
    df = get_data("ASISTENCIAS")
    if df.empty or "fecha" not in df.columns:
        return {}
    df_filtro = df[(df["grupo_id"] == grupo_id) & (df["fecha"] == str(fecha))]
    return {str(row["matricula"]): bool(row["presente"]) for _, row in df_filtro.iterrows()}

# ==========================================
# REPORTES EN PDF Y EXCEL
# ==========================================
def generar_reporte_pdf(grupo_id, nombre_dip, grupo, docente):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter,
        rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )
    story = []
    styles = getSampleStyleSheet()

    logo_izq = RLImage(PATH_LOGO_UADEC, width=70, height=70) if os.path.exists(PATH_LOGO_UADEC) else ""
    logo_der = RLImage(PATH_LOGO_POSGRADO, width=70, height=70) if os.path.exists(PATH_LOGO_POSGRADO) else ""

    header_text = Paragraph(
        f"<font size=11><b>UNIVERSIDAD AUTÓNOMA DE COAHUILA</b></font><br/>"
        f"<font size=9 color='#2D3748'>FACULTAD DE INGENIERÍA - COORDINACIÓN DE POSGRADO</font><br/>"
        f"<font size=8 color='#718096'>REPORTE OFICIAL DE PARTICIPANTES</font>",
        ParagraphStyle('HeaderTxt', alignment=1, leading=13)
    )

    t_header = Table([[logo_izq, header_text, logo_der]], colWidths=[80, 380, 80])
    t_header.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (0,0), 'LEFT'),
        ('ALIGN', (2,0), (2,0), 'RIGHT'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(t_header)
    story.append(Spacer(1, 10))

    meta_style = ParagraphStyle('Meta', parent=styles['Normal'], fontSize=8.5, leading=12)
    meta_data = [
        [
            Paragraph(f"<b>DIPLOMADO:</b> {nombre_dip.upper()}", meta_style),
            Paragraph(f"<b>GRUPO / GENERACIÓN:</b> {grupo.upper()}", meta_style)
        ],
        [
            Paragraph(f"<b>INSTRUCTOR:</b> {docente.upper()}", meta_style),
            Paragraph(f"<b>FECHA DE EMISIÓN:</b> {date.today().strftime('%d/%m/%Y')}", meta_style)
        ]
    ]
    t_meta = Table(meta_data, colWidths=[270, 270])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#EDF2F7")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E0")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 15))

    cell_head = ParagraphStyle('CH', parent=styles['Normal'], fontSize=8.5, fontName='Helvetica-Bold', textColor=colors.whitesmoke, alignment=1)
    cell_body = ParagraphStyle('CB', parent=styles['Normal'], fontSize=8, leading=10)

    data = [[
        Paragraph("#", cell_head),
        Paragraph("MATRÍCULA", cell_head),
        Paragraph("NOMBRE COMPLETO", cell_head),
        Paragraph("CARRERA / ESPECIALIDAD", cell_head),
        Paragraph("TELÉFONO", cell_head)
    ]]

    df_al = get_alumnos_df(grupo_id)
    for idx, row in df_al.reset_index(drop=True).iterrows():
        data.append([
            Paragraph(str(idx + 1), cell_body),
            Paragraph(str(row["matricula"]), cell_body),
            Paragraph(str(row["nombre"]), cell_body),
            Paragraph(str(row["carrera"]), cell_body),
            Paragraph(str(row["telefono"]), cell_body)
        ])

    t_alumnos = Table(data, colWidths=[25, 75, 210, 150, 80], repeatRows=1)
    t_alumnos.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1A365D")),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_alumnos)
    doc.build(story)
    buffer.seek(0)
    return buffer

def generar_excel(grupo_id, nom_dip, nom_gpo, nom_doc, con_pagos=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "LISTA OFICIAL"
    ws.views.sheetView[0].showGridLines = True

    df_al = get_alumnos_df(grupo_id)
    col_fin_letra = "H" if con_pagos else "E"
    col_fin_num = 8 if con_pagos else 5

    if OpenpyxlImage is not None:
        if os.path.exists(PATH_LOGO_UADEC):
            try:
                img_uadec = OpenpyxlImage(PATH_LOGO_UADEC)
                img_uadec.width, img_uadec.height = 75, 75
                ws.add_image(img_uadec, "A1")
            except Exception:
                pass
        if os.path.exists(PATH_LOGO_POSGRADO):
            try:
                img_pos = OpenpyxlImage(PATH_LOGO_POSGRADO)
                img_pos.width, img_pos.height = 75, 75
                ws.add_image(img_pos, f"{col_fin_letra}1")
            except Exception:
                pass

    col_centro = chr(64 + col_fin_num - 1)
    ws.merge_cells(f"B1:{col_centro}1")
    ws["B1"] = "UNIVERSIDAD AUTÓNOMA DE COAHUILA"
    ws["B1"].font = Font(name="Calibri", size=13, bold=True, color="1A365D")
    ws["B1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"B2:{col_centro}2")
    ws["B2"] = "FACULTAD DE INGENIERÍA - COORDINACIÓN DE POSGRADO"
    ws["B2"].font = Font(name="Calibri", size=10, bold=True, color="4A5568")
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"B3:{col_centro}3")
    ws["B3"] = "REPORTE OFICIAL DE PARTICIPANTES" if not con_pagos else "CONTROL INTERNO Y REPORTE DE PAGOS"
    ws["B3"].font = Font(name="Calibri", size=9, bold=False, color="718096")
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height = 24, 18, 16

    fill_meta = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    font_meta = Font(name="Calibri", size=9, color="1A202C")
    thin_border = Border(left=Side(style='thin', color="CBD5E0"), right=Side(style='thin', color="CBD5E0"),
                         top=Side(style='thin', color="CBD5E0"), bottom=Side(style='thin', color="CBD5E0"))

    ws["A5"], ws["D5"] = f"DIPLOMADO: {nom_dip}", f"GRUPO / GENERACIÓN: {nom_gpo}"
    ws["A6"], ws["D6"] = f"INSTRUCTOR: {nom_doc}", f"FECHA DE EMISIÓN: {date.today().strftime('%d/%m/%Y')}"
    ws["A5"].font, ws["D5"].font, ws["A6"].font, ws["D6"].font = font_meta, font_meta, font_meta, font_meta

    ws.merge_cells("A5:C5"); ws.merge_cells(f"D5:{col_fin_letra}5")
    ws.merge_cells("A6:C6"); ws.merge_cells(f"D6:{col_fin_letra}6")

    for r in range(5, 7):
        ws.row_dimensions[r].height = 20
        for col_idx in range(1, col_fin_num + 1):
            cell = ws.cell(row=r, column=col_idx)
            cell.fill, cell.border, cell.alignment = fill_meta, thin_border, Alignment(vertical="center", indent=1)

    headers = ["#", "MATRÍCULA", "NOMBRE COMPLETO", "CARRERA / ESPECIALIDAD", "TELÉFONO"]
    if con_pagos:
        headers.extend(["1ER PAGO", "2DO PAGO", "PAGO COMPLETO"])

    ws.row_dimensions[8].height = 24
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
    for c_idx, h_text in enumerate(headers, 1):
        c = ws.cell(row=8, column=c_idx, value=h_text)
        c.fill, c.font, c.alignment = header_fill, header_font, Alignment(horizontal="center", vertical="center")

    body_font = Font(name="Calibri", size=9, color="2D3748")
    zebra_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    cell_border = Border(left=Side(style='thin', color="E2E8F0"), right=Side(style='thin', color="E2E8F0"),
                         top=Side(style='thin', color="E2E8F0"), bottom=Side(style='thin', color="E2E8F0"))

    fila_actual = 9
    for idx, row in df_al.reset_index(drop=True).iterrows():
        ws.row_dimensions[fila_actual].height = 20
        fill_row = zebra_fill if idx % 2 == 1 else white_fill
        valores = [idx + 1, row["matricula"], row["nombre"], row["carrera"], row["telefono"]]
        if con_pagos:
            valores.extend([
                "PAGADO" if row["pago_1"] else "PENDIENTE",
                "PAGADO" if row["pago_2"] else "PENDIENTE",
                "LIQUIDADO" if row["pago_completo"] else "PENDIENTE"
            ])
        for c_idx, val in enumerate(valores, 1):
            cell = ws.cell(row=fila_actual, column=c_idx, value=val)
            cell.font, cell.fill, cell.border = body_font, fill_row, cell_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 2] or (con_pagos and c_idx >= 6) else "left", vertical="center", indent=1)
        fila_actual += 1

    anchos = {"A": 6, "B": 15, "C": 38, "D": 32, "E": 16}
    if con_pagos:
        anchos["F"], anchos["G"], anchos["H"] = 14, 14, 16
    for col_letter, width in anchos.items():
        ws.column_dimensions[col_letter].width = width

    if con_pagos:
        df_h = get_horas_df(grupo_id)
        if not df_h.empty:
            ws_h = wb.create_sheet(title="BITÁCORA DE HORAS")
            ws_h.views.sheetView[0].showGridLines = True
            for c_idx, h_text in enumerate(["FECHA", "HORAS IMPARTIDAS", "TEMA / ACTIVIDAD ABORDADA"], 1):
                c = ws_h.cell(row=1, column=c_idx, value=h_text)
                c.fill, c.font, c.alignment = header_fill, header_font, Alignment(horizontal="center", vertical="center")
            for h_idx, h_row in df_h.reset_index(drop=True).iterrows():
                f_idx = h_idx + 2
                for col_num, val in enumerate([str(h_row["fecha"]), float(h_row["horas"]), str(h_row["tema"])], 1):
                    c = ws_h.cell(row=f_idx, column=col_num, value=val)
                    c.font, c.fill, c.border = body_font, (zebra_fill if h_idx % 2 == 1 else white_fill), cell_border
                    c.alignment = Alignment(horizontal="center" if col_num in [1, 2] else "left", vertical="center", indent=1)
            ws_h.column_dimensions["A"].width, ws_h.column_dimensions["B"].width, ws_h.column_dimensions["C"].width = 15, 18, 50

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ==========================================
# INTERFAZ DE USUARIO (OPTIMIZADA PARA MÓVIL)
# ==========================================
col_l1, col_t, col_l2 = st.columns([1, 4, 1])
with col_l1:
    if os.path.exists(PATH_LOGO_UADEC):
        st.image(PATH_LOGO_UADEC, use_container_width=True)
    else:
        st.caption("🏛️️ UAdeC")
with col_t:
    st.markdown("<h2 style='text-align: center; margin-bottom: 0;'>COORDINACIÓN DE POSGRADO DE INGENIERÍA</h2>", unsafe_allow_html=True)
    st.markdown("<h4 style='text-align: center; color: gray; margin-top: 0;'>SISTEMA DE GESTIÓN Y CONTROL DE DIPLOMADOS</h4>", unsafe_allow_html=True)
with col_l2:
    if os.path.exists(PATH_LOGO_POSGRADO):
        st.image(PATH_LOGO_POSGRADO, use_container_width=True)
    else:
        st.caption("🔬 POSGRADO")

st.divider()

df_grupos = get_grupos()

if df_grupos.empty:
    st.info("👋 NO HAY NINGÚN GRUPO REGISTRADO EN GOOGLE SHEETS.")
    with st.form("form_primer_grupo"):
        nvo_dip = st.text_input("NOMBRE DEL DIPLOMADO").upper()
        nvo_gpo = st.text_input("GRUPO / CLAVE (EJ. GPO. 2026-I)").upper()
        nvo_doc = st.text_input("DOCENTE RESPONSABLE").upper()
        if st.form_submit_button("REGISTRAR Y GUARDAR EN NUBE"):
            if nvo_dip and nvo_gpo:
                add_grupo(nvo_dip, nvo_gpo, nvo_doc)
                st.success("GRUPO GUARDADO EN GOOGLE SHEETS.")
                st.rerun()
            else:
                st.error("TODOS LOS CAMPOS SON OBLIGATORIOS.")
    st.stop()

# Selector visible y amigable en móvil
col_sel, col_btn_nvo = st.columns([3, 1])
opciones_grupos = {row["id"]: f"{row['grupo']} - {row['diplomado']}" for _, row in df_grupos.iterrows()}

with col_sel:
    grupo_seleccionado_id = st.selectbox(
        "SELECCIONA EL DIPLOMADO / GRUPO",
        options=list(opciones_grupos.keys()),
        format_func=lambda x: opciones_grupos[x],
        key="selector_grupo"
    )

info_grupo = df_grupos[df_grupos["id"] == grupo_seleccionado_id].iloc[0]
dip_nombre, dip_grupo, dip_docente = info_grupo["diplomado"], info_grupo["grupo"], info_grupo["docente"]

with col_btn_nvo:
    st.write("")
    st.write("")
    with st.popover("➕ NUEVO GRUPO"):
        with st.form("form_nuevo_grupo_movil"):
            nvo_dip = st.text_input("NOMBRE DEL DIPLOMADO").upper()
            nvo_gpo = st.text_input("GRUPO / CLAVE").upper()
            nvo_doc = st.text_input("DOCENTE RESPONSABLE", value=dip_docente).upper()
            if st.form_submit_button("CREAR"):
                if nvo_dip and nvo_gpo:
                    add_grupo(nvo_dip, nvo_gpo, nvo_doc)
                    st.success("GRUPO GUARDADO.")
                    st.rerun()

with st.sidebar:
    st.header("⚙️ OPCIONES")
    st.info(f"**Diplomado:** {dip_nombre}\n\n**Grupo:** {dip_grupo}\n\n**Docente:** {dip_docente}")
    df_horas_tot = get_horas_df(grupo_seleccionado_id)
    tot_h = df_horas_tot["horas"].sum() if not df_horas_tot.empty else 0.0
    st.metric("HORAS ACUMULADAS", f"{tot_h:.1f} HRS")

    st.divider()
    st.subheader("📥 DESCARGAS")
    pdf_bytes = generar_reporte_pdf(grupo_seleccionado_id, dip_nombre, dip_grupo, dip_docente)
    st.download_button("📄 PDF OFICIAL (SIN PAGOS)", pdf_bytes, f"LISTA_OFICIAL_{dip_grupo}.pdf", "application/pdf", use_container_width=True)

    excel_sin = generar_excel(grupo_seleccionado_id, dip_nombre, dip_grupo, dip_docente, con_pagos=False)
    st.download_button("📊 EXCEL OFICIAL (SIN PAGOS)", excel_sin, f"ALUMNOS_{dip_grupo}_OFICIAL.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)

    excel_con = generar_excel(grupo_seleccionado_id, dip_nombre, dip_grupo, dip_docente, con_pagos=True)
    st.download_button("💰 EXCEL (CON PAGOS Y HORAS)", excel_con, f"CONTROL_{dip_grupo}_PAGOS.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)

    with st.expander("🗑️ ELIMINAR GRUPO"):
        if st.button("CONFIRMAR Y BORRAR ESTE GRUPO"):
            delete_grupo(grupo_seleccionado_id)
            st.warning("GRUPO ELIMINADO.")
            st.rerun()

# Pestañas principales
tab_alumnos, tab_asistencia, tab_horas = st.tabs([
    f"📋 ALUMNOS Y PAGOS ({dip_grupo})",
    "✅ PASE DE LISTA",
    "⏱️ BITÁCORA DE HORAS"
])

with tab_alumnos:
    df_al = get_alumnos_df(grupo_seleccionado_id)
    if not df_al.empty:
        df_display = df_al.rename(columns={
            "matricula": "MATRÍCULA", "nombre": "NOMBRE", "carrera": "CARRERA",
            "telefono": "TELÉFONO", "pago_1": "1ER PAGO", "pago_2": "2DO PAGO", "pago_completo": "PAGO COMPLETO"
        })
        edited_data = st.data_editor(
            df_display, use_container_width=True, num_rows="fixed",
            column_config={
                "MATRÍCULA": st.column_config.TextColumn(required=True),
                "NOMBRE": st.column_config.TextColumn(required=True),
                "1ER PAGO": st.column_config.CheckboxColumn(),
                "2DO PAGO": st.column_config.CheckboxColumn(),
                "PAGO COMPLETO": st.column_config.CheckboxColumn()
            },
            key=f"editor_{grupo_seleccionado_id}"
        )
        c_save, c_del = st.columns(2)
        if c_save.button("💾 GUARDAR CAMBIOS EN LA NUBE", use_container_width=True):
            for _, r in edited_data.iterrows():
                save_alumno(r["MATRÍCULA"], grupo_seleccionado_id, r["NOMBRE"], r["CARRERA"], r["TELÉFONO"], r["1ER PAGO"], r["2DO PAGO"], r["PAGO COMPLETO"])
            st.success("DATOS GUARDADOS PERMANENTEMENTE EN GOOGLE SHEETS.")
            st.rerun()
        with c_del:
            with st.popover("🗑️ ELIMINAR UN ALUMNO"):
                mat_del = st.selectbox("MATRÍCULA", df_al["matricula"].tolist())
                if st.button("CONFIRMAR"):
                    delete_alumno(mat_del, grupo_seleccionado_id)
                    st.rerun()
    else:
        st.info("NO HAY ALUMNOS EN ESTE GRUPO.")

    with st.expander("➕ REGISTRAR NUEVO ALUMNO", expanded=df_al.empty):
        with st.form("form_nuevo_alumno"):
            c1, c2 = st.columns(2)
            n_mat, n_nom = c1.text_input("MATRÍCULA").upper(), c2.text_input("NOMBRE COMPLETO").upper()
            c3, c4 = st.columns(2)
            n_car, n_tel = c3.text_input("CARRERA").upper(), c4.text_input("TELÉFONO").upper()
            cp1, cp2, cp3 = st.columns(3)
            p1, p2, p_comp = cp1.checkbox("1ER PAGO"), cp2.checkbox("2DO PAGO"), cp3.checkbox("PAGO COMPLETO")
            if st.form_submit_button("GUARDAR EN GOOGLE SHEETS"):
                if n_mat and n_nom:
                    save_alumno(n_mat, grupo_seleccionado_id, n_nom, n_car, n_tel, p1, p2, p_comp)
                    st.success("ALUMNO GUARDADO.")
                    st.rerun()
                else:
                    st.error("MATRÍCULA Y NOMBRE REQUERIDOS.")

with tab_asistencia:
    f_asis = st.date_input("FECHA DE LA SESIÓN", value=date.today())
    str_f = str(f_asis)
    asis_guardadas = get_asistencia_fecha(grupo_seleccionado_id, str_f)
    df_al_asis = get_alumnos_df(grupo_seleccionado_id)

    if not df_al_asis.empty:
        with st.form("form_asistencia"):
            estados = {}
            for _, al in df_al_asis.iterrows():
                val_prev = asis_guardadas.get(str(al["matricula"]), True)
                cn, cc = st.columns([3, 1])
                cn.write(f"**{al['nombre']}** ({al['matricula']})")
                estados[al["matricula"]] = cc.checkbox("PRESENTE", value=val_prev, key=f"att_{str_f}_{al['matricula']}")
            if st.form_submit_button("💾 GUARDAR ASISTENCIAS EN GOOGLE SHEETS"):
                save_asistencia_fecha(grupo_seleccionado_id, str_f, estados)
                st.success("ASISTENCIA GUARDADA.")
    else:
        st.info("REGISTRE ALUMNOS PRIMERO.")

with tab_horas:
    c_h1, c_h2 = st.columns(2)
    with c_h1:
        with st.form("form_horas"):
            f_h = st.date_input("FECHA", value=date.today())
            n_hrs = st.number_input("HORAS DADAS", min_value=0.5, max_value=12.0, value=2.0, step=0.5)
            n_tema = st.text_input("TEMA / CONTENIDO ABORDADO").upper()
            if st.form_submit_button("GUARDAR HORAS EN NUBE"):
                if n_tema:
                    save_hora(grupo_seleccionado_id, f_h, n_hrs, n_tema)
                    st.success("HORAS REGISTRADAS EN GOOGLE SHEETS.")
                    st.rerun()
    with c_h2:
        df_hist = get_horas_df(grupo_seleccionado_id)
        if not df_hist.empty:
            st.dataframe(df_hist[["fecha", "horas", "tema"]].rename(columns={"fecha": "FECHA", "horas": "HORAS", "tema": "TEMA"}), use_container_width=True)
        else:
            st.caption("NO HAY HORAS REGISTRADAS.")
